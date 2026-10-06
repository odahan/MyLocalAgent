from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


def clean(text: str) -> str:
    return " ".join(text.split())


path = Path(sys.argv[1])
doc = Document(path)
outline = []
images = []
current = {"h1": "", "h2": "", "h3": ""}
paragraphs = list(doc.paragraphs)

for index, paragraph in enumerate(paragraphs):
    text = clean(paragraph.text)
    style = paragraph.style.name if paragraph.style else ""
    if style in {"Heading 1", "Heading 2", "Heading 3"}:
        level = style[-1]
        current[f"h{level}"] = text
        if level == "1":
            current["h2"] = ""
            current["h3"] = ""
        elif level == "2":
            current["h3"] = ""
        outline.append({"index": index, "level": int(level), "text": text})

    drawings = paragraph._p.xpath(".//w:drawing")
    if drawings:
        before = []
        after = []
        for p in paragraphs[max(0, index - 3):index]:
            value = clean(p.text)
            if value:
                before.append(value)
        for p in paragraphs[index + 1:index + 4]:
            value = clean(p.text)
            if value:
                after.append(value)
        doc_pr = paragraph._p.xpath(".//wp:docPr")
        meta = []
        for node in doc_pr:
            meta.append({
                "name": node.get("name", ""),
                "descr": node.get("descr", ""),
                "title": node.get("title", ""),
            })
        images.append({
            "paragraph_index": index,
            "headings": dict(current),
            "paragraph_text": text,
            "before": before,
            "after": after,
            "metadata": meta,
        })

result = {
    "path": str(path),
    "paragraph_count": len(paragraphs),
    "table_count": len(doc.tables),
    "section_count": len(doc.sections),
    "inline_shape_count": len(doc.inline_shapes),
    "outline": outline,
    "images": images,
}
if len(sys.argv) > 2 and sys.argv[2] == "compact":
    chapter_openings = []
    for heading in outline:
        if heading["level"] != 1 or heading["index"] < 70:
            continue
        start = heading["index"]
        body = []
        for p in paragraphs[start + 1:]:
            style = p.style.name if p.style else ""
            if style in {"Heading 1", "Heading 2"}:
                break
            value = clean(p.text)
            if value:
                body.append(value)
            if len(body) >= 7:
                break
        chapter_openings.append({"heading": heading["text"], "body": body})
    compact_images = []
    for item in images:
        caption = next((t for t in item["after"] if t.startswith("Figure ")), "")
        compact_images.append({
            "paragraph_index": item["paragraph_index"],
            "chapter": item["headings"]["h1"],
            "section": item["headings"]["h2"],
            "caption": caption,
        })
    print(json.dumps({"chapter_openings": chapter_openings, "images": compact_images}, ensure_ascii=False, indent=2))
else:
    print(json.dumps(result, ensure_ascii=False, indent=2))
