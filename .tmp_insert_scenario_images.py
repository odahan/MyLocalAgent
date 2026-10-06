from __future__ import annotations

import copy
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches


SOURCE = Path(sys.argv[1])
OUTPUT = Path(sys.argv[2])
IMAGES = Path(sys.argv[3])


def normalized(text: str) -> str:
    return " ".join(text.split())


doc = Document(SOURCE)

# Shift the existing captions before introducing the new figures.
renumber = {
    "Figure 1.1": "Figure 1.2",
    "Figure 1.2": "Figure 1.3",
    "Figure 2.1": "Figure 2.2",
    "Figure 2.2": "Figure 2.3",
    "Figure 3.1": "Figure 3.2",
    "Figure 3.2": "Figure 3.3",
    "Figure 4.1": "Figure 4.2",
    "Figure 4.2": "Figure 4.3",
    "Figure 5.1": "Figure 5.2",
    "Figure 5.2": "Figure 5.3",
    "Figure 6.1": "Figure 6.2",
    "Figure 6.2": "Figure 6.3",
    "Figure 7.1": "Figure 7.2",
    "Figure 7.2": "Figure 7.3",
    "Figure 8.1": "Figure 8.2",
    "Figure 8.2": "Figure 8.3",
    "Figure I.1": "Figure I.2",
    "Figure 9.1": "Figure 9.2",
    "Figure 9.2": "Figure 9.4",
    "Figure 9.3": "Figure 9.7",
}

for paragraph in doc.paragraphs:
    text = paragraph.text
    if not text.startswith("Figure "):
        continue
    for old, new in renumber.items():
        if text.startswith(old + " "):
            paragraph.text = new + text[len(old):]
            break

# Copy the paragraph properties used by the existing inline figures and captions.
image_template = next(p for p in doc.paragraphs if p._p.xpath(".//w:drawing"))
caption_template = next(p for p in doc.paragraphs if p.style and p.style.name == "Caption" and p.text.startswith("Figure "))


def find_anchor(starts_with: str) -> Paragraph:
    matches = [p for p in doc.paragraphs if normalized(p.text).startswith(starts_with)]
    if len(matches) != 1:
        raise RuntimeError(f"Anchor {starts_with!r}: expected 1 match, found {len(matches)}")
    return matches[0]


def insert_after(anchor: Paragraph, image_name: str, caption: str) -> None:
    image_path = IMAGES / image_name
    if not image_path.exists():
        raise FileNotFoundError(image_path)

    image_xml = OxmlElement("w:p")
    if image_template._p.pPr is not None:
        image_xml.append(copy.deepcopy(image_template._p.pPr))
    anchor._p.addnext(image_xml)
    image_paragraph = Paragraph(image_xml, anchor._parent)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.add_run().add_picture(str(image_path), width=Inches(4.84))

    caption_xml = OxmlElement("w:p")
    if caption_template._p.pPr is not None:
        caption_xml.append(copy.deepcopy(caption_template._p.pPr))
    image_xml.addnext(caption_xml)
    caption_paragraph = Paragraph(caption_xml, anchor._parent)
    caption_paragraph.style = caption_template.style
    caption_paragraph.add_run(caption)


# Insert from the end of the document toward the beginning so XML positions stay predictable.
items = [
    (
        "Une même organisation envisage deux changements.",
        "09D_Scene_tri_et_acces_assistance.png",
        "Figure 9.6 - Scénario construit. Une préférence de tri locale et un accès temporaire aux dossiers d'une autre organisation n'engagent ni les mêmes preuves ni les mêmes autorités.",
    ),
    (
        "La tentation est de transformer ce prototype en fondation durable",
        "09C_Scene_prototype_espace_partage.png",
        "Figure 9.5 - Scénario construit. Le prototype fonctionne pour deux organisations, mais son espace de données commun laisse apparaître les réservations de l'une dans l'interface de l'autre.",
    ),
    (
        "Cette facilité est précisément ce qui rend l'arbitrage risqué.",
        "09B_Scene_enquete_references.png",
        "Figure 9.3 - Scénario construit. Le support, le métier et la maintenance réunissent des fragments de connaissance avant de qualifier les références ambiguës.",
    ),
    (
        "Le responsable produit réduit donc le périmètre de la livraison.",
        "09A_Scene_demande_planificateur.png",
        "Figure 9.1 - Scénario construit. Le client voit que sa demande a été reçue tandis que le nouveau créneau reste en attente de la décision du planificateur.",
    ),
    (
        "Dans une boîte à outils comme Spec Kit",
        "08bis_Scene_dossier_SDD.png",
        "Figure I.1 - Cadre pédagogique. Le dossier SDD relie les artefacts de travail au cas tarifaire sans faire disparaître la règle encore inconnue.",
    ),
    (
        "Le paiement lui-même n'est pas rejoué dans ce scénario.",
        "08_Scene_notifications_dupliquees.png",
        "Figure 8.1 - Scénario construit. Une relance automatique peut délivrer le message attendu et produire un doublon visible par son destinataire.",
    ),
    (
        "Un consommateur ancien valide strictement les champs qu'il accepte",
        "07_Scene_evenement_rejete.png",
        "Figure 7.1 - Scénario construit. Le producteur émet un événement valide localement que l'ancien consommateur rejette.",
    ),
    (
        "Deux formes de code peuvent alors sembler naturelles.",
        "06_Scene_tarif_livraison_urgente.png",
        "Figure 6.1 - Scénario construit. Le montant de l'option urgente dépend du prix, de l'heure reçue, du créneau confirmé et de l'état de préparation.",
    ),
    (
        "Les journaux bruts sont copiés sans distinction dans un index de recherche",
        "05_Scene_console_incident.png",
        "Figure 5.1 - Scénario construit. La console réunit les informations utiles à l'astreinte, mais expose aussi des données qui n'étaient pas destinées à cette consultation.",
    ),
    (
        "Lors d'un échange, l'opérateur formule ce qui lui manque",
        "04_Scene_import_partiel.png",
        "Figure 4.1 - Scénario construit. Le tableau de bord montre un lot partiellement traité sans établir si sa relance rejouera les lignes déjà appliquées.",
    ),
    (
        "Mais un commentaire ancien attire l'attention",
        "03_Scene_collision_references_client.png",
        "Figure 3.1 - Scénario construit. Deux références distinctes produisent la même clé avant que leur équivalence métier soit établie.",
    ),
    (
        "Le tableau de suivi affiche donc deux histoires.",
        "02_Scene_file_pull_requests.png",
        "Figure 2.1 - Scénario construit. Les pull requests s'accumulent plus vite que les mainteneurs ne peuvent relier chaque changement aux usages du kit partagé.",
    ),
    (
        "La première modifie directement le créneau dans la commande.",
        "01_Scene_modification_creneau.png",
        "Figure 1.1 - Scénario construit. Le client formule une demande de changement tandis que le créneau initial reste confirmé et que le planificateur doit encore traiter la demande.",
    ),
]

for anchor_text, image_name, caption in items:
    insert_after(find_anchor(anchor_text), image_name, caption)

doc.save(OUTPUT)
print(f"Saved {OUTPUT}")
